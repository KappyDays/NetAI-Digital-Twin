"""Generate architecture diagram as PowerPoint (pptx) with editable shapes.

Usage:
    pip install python-pptx
    python generate_pptx.py

Output:
    docs/brandt_architecture.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os


def hex_to_rgb(hex_str):
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def add_box(slide, left, top, width, height, text, subtext="",
            fill="#E8F4FD", border="#2196F3", font_size=12, sub_size=9,
            bold=True, font_color="#1a1a1a"):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_to_rgb(fill)
    shape.line.color.rgb = hex_to_rgb(border)
    shape.line.width = Pt(1.5)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = hex_to_rgb(font_color)

    if subtext:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = subtext
        run2.font.size = Pt(sub_size)
        run2.font.color.rgb = hex_to_rgb("#555555")
        run2.font.italic = True

    return shape


def add_text(slide, left, top, width, height, text, font_size=9,
             bold=False, color="#333333", align=PP_ALIGN.LEFT, italic=False):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = hex_to_rgb(color)
    run.font.italic = italic
    return txBox


def add_arrow(slide, start_x, start_y, end_x, end_y, color="#555555", width=1.5):
    connector = slide.shapes.add_connector(
        1,  # straight connector
        start_x, start_y, end_x, end_y
    )
    connector.line.color.rgb = hex_to_rgb(color)
    connector.line.width = Pt(width)
    return connector


def I(val):
    return Inches(val)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    # ── Title ──
    add_text(slide, I(0.5), I(0.15), I(12), I(0.5),
             "BranDT: OpenUSD Digital Twin Data Architecture",
             font_size=22, bold=True, color="#1a1a1a", align=PP_ALIGN.CENTER)
    add_text(slide, I(0.5), I(0.55), I(12), I(0.3),
             "Nucleus -> Iceberg Lakehouse -> Time Travel Pipeline",
             font_size=11, color="#666666", align=PP_ALIGN.CENTER, italic=True)

    # ══════════════════════════════════════════════════════════
    # Layer 1: Isaac Sim + Nucleus
    # ══════════════════════════════════════════════════════════

    # Isaac Sim
    add_box(slide, I(0.3), I(1.0), I(4.5), I(1.6),
            "NVIDIA Isaac Sim",
            "USD Stage  |  Viewport  |  Extension",
            fill="#FFF3E0", border="#FF9800", font_size=14, sub_size=9)

    add_text(slide, I(0.6), I(1.65), I(2.5), I(0.7),
             "/World\n  /Robots/Jetbot   (ref: jetbot.usd)\n  /Environment/Table (ref: table.usd)\n  /Props/Block_A    (ref: block.usd)",
             font_size=7, color="#555555")

    # Reference / Payload badges
    add_box(slide, I(3.5), I(1.65), I(0.7), I(0.25),
            "Reference", fill="#FFCCBC", border="#E64A19",
            font_size=7, bold=True, font_color="#BF360C")
    add_box(slide, I(4.25), I(1.65), I(0.55), I(0.25),
            "Payload", fill="#E1BEE7", border="#7B1FA2",
            font_size=7, bold=True, font_color="#4A148C")

    # Nucleus Server
    add_box(slide, I(5.8), I(1.0), I(3.8), I(1.6),
            "Omniverse Nucleus",
            "USD File Server (omniverse://)",
            fill="#F3E5F5", border="#9C27B0", font_size=14, sub_size=9)

    add_text(slide, I(6.1), I(1.65), I(3), I(0.7),
             "scene.usd  (root layer)\njetbot.usd  |  kaya.usd  |  table.usd\nbasic_block.usd  |  grid.usd",
             font_size=7, color="#555555")

    # Arrow: Isaac Sim <-> Nucleus
    add_text(slide, I(4.85), I(1.3), I(0.9), I(0.25),
             "Save / Open", font_size=8, color="#9C27B0",
             align=PP_ALIGN.CENTER, bold=True)
    add_arrow(slide, I(4.8), I(1.7), I(5.8), I(1.7), color="#9C27B0", width=2)

    # ══════════════════════════════════════════════════════════
    # Layer 2: Nucleus Pipeline
    # ══════════════════════════════════════════════════════════

    # Pipeline container
    add_box(slide, I(0.3), I(3.0), I(9.3), I(2.2),
            "", fill="#F1F8E9", border="#4CAF50", font_size=1)

    add_text(slide, I(0.5), I(3.05), I(9), I(0.3),
             "Nucleus Pipeline  (nucleus_pipeline/)",
             font_size=12, bold=True, color="#2E7D32", align=PP_ALIGN.CENTER)
    add_text(slide, I(0.5), I(3.3), I(9), I(0.2),
             "Python CLI  |  PyUSD (usd-core)  |  omniverseclient  |  No Isaac Sim Required",
             font_size=8, color="#558B2F", align=PP_ALIGN.CENTER, italic=True)

    # Step boxes
    step_y = I(3.6)
    step_h = I(1.3)

    add_box(slide, I(0.5), step_y, I(2.0), step_h,
            "1. Download",
            "omniverseclient\n(subprocess isolation)",
            fill="#E8F5E9", border="#4CAF50", font_size=11, sub_size=7)

    add_box(slide, I(2.7), step_y, I(2.2), step_h,
            "2. Parse USD",
            "Entity Detection\nOverride Extraction\nAsset URL Parsing",
            fill="#E8F5E9", border="#4CAF50", font_size=11, sub_size=7)

    add_box(slide, I(5.1), step_y, I(2.2), step_h,
            "3. Generate",
            "Create root.usda\nRewrite Ref/Payload\npaths to ./entities/*",
            fill="#E8F5E9", border="#4CAF50", font_size=11, sub_size=7)

    add_box(slide, I(7.5), step_y, I(1.9), step_h,
            "4. Store",
            "MinIO Upload\nIceberg INSERT",
            fill="#E8F5E9", border="#4CAF50", font_size=11, sub_size=7)

    # Step arrows
    add_arrow(slide, I(2.5), I(4.25), I(2.7), I(4.25), color="#4CAF50", width=1.5)
    add_arrow(slide, I(4.9), I(4.25), I(5.1), I(4.25), color="#4CAF50", width=1.5)
    add_arrow(slide, I(7.3), I(4.25), I(7.5), I(4.25), color="#4CAF50", width=1.5)

    # Arrow: Nucleus -> Pipeline
    add_text(slide, I(4.0), I(2.7), I(2.5), I(0.2),
             "omni.client.read_file()", font_size=7, color="#9C27B0",
             align=PP_ALIGN.CENTER, italic=True)
    add_arrow(slide, I(7.7), I(2.6), I(1.5), I(3.0), color="#9C27B0", width=1.5)

    # ══════════════════════════════════════════════════════════
    # Layer 3: Storage
    # ══════════════════════════════════════════════════════════

    # MinIO
    add_box(slide, I(0.3), I(5.6), I(4.2), I(1.7),
            "MinIO (S3-Compatible Storage)",
            "", fill="#FFF8E1", border="#FFA000", font_size=11)

    add_text(slide, I(0.5), I(6.0), I(4), I(1.2),
             "backups/{timestamp}/\n"
             "  root.usda        (full Stage local data)\n"
             "  entities/\n"
             "    jetbot.usd      (original asset)\n"
             "    kaya.usd\n"
             "    basic_block.usd  (shared, deduplicated)",
             font_size=7, color="#555555")

    # Iceberg
    add_box(slide, I(4.8), I(5.6), I(4.8), I(1.7),
            "Apache Iceberg (via Trino SQL)",
            "", fill="#E8EAF6", border="#3F51B5", font_size=11)

    add_text(slide, I(5.0), I(6.0), I(4.5), I(1.2),
             "entities\n"
             "  entity_path | type | hash | backup_source | time\n\n"
             "prim_snapshots\n"
             "  entity_path | properties (Override JSON) | hash\n\n"
             "Time Travel: query any past state via SQL",
             font_size=7, color="#333355")

    # Arrows: Pipeline -> Storage
    add_text(slide, I(3.5), I(5.25), I(1.5), I(0.2),
             "USD Files", font_size=7, color="#FFA000", bold=True, align=PP_ALIGN.CENTER)
    add_arrow(slide, I(8.0), I(5.2), I(2.4), I(5.6), color="#FFA000", width=1.5)

    add_text(slide, I(7.0), I(5.25), I(1.8), I(0.2),
             "Override + Metadata", font_size=7, color="#3F51B5", bold=True, align=PP_ALIGN.CENTER)
    add_arrow(slide, I(8.5), I(5.2), I(7.2), I(5.6), color="#3F51B5", width=1.5)

    # ══════════════════════════════════════════════════════════
    # Right side: Time Travel + Dashboard
    # ══════════════════════════════════════════════════════════

    # Time Travel
    add_box(slide, I(10.3), I(1.5), I(2.7), I(3.0),
            "Time Travel", "",
            fill="#E0F7FA", border="#00ACC1", font_size=14)

    add_text(slide, I(10.5), I(2.2), I(2.3), I(0.5),
             "Query", font_size=11, bold=True, color="#006064",
             align=PP_ALIGN.CENTER)
    add_text(slide, I(10.5), I(2.6), I(2.3), I(0.5),
             "\"Where was Jetbot\n2 hours ago?\"",
             font_size=9, color="#00695C", align=PP_ALIGN.CENTER, italic=True)

    add_text(slide, I(10.5), I(3.2), I(2.3), I(0.3),
             "Restore", font_size=11, bold=True, color="#006064",
             align=PP_ALIGN.CENTER)
    add_text(slide, I(10.5), I(3.5), I(2.3), I(0.8),
             "Download root.usda from MinIO\n-> Open in Isaac Sim\n-> Past Stage restored",
             font_size=8, color="#00695C", align=PP_ALIGN.CENTER)

    # Dashboard
    add_box(slide, I(10.3), I(5.6), I(2.7), I(1.7),
            "Web Dashboard",
            "Entity Diff  |  SQL Query\nCongestion  |  Iceberg Hub\nReact :3000",
            fill="#FCE4EC", border="#E91E63", font_size=12, sub_size=8)

    # Arrow: Iceberg -> Time Travel
    add_arrow(slide, I(9.6), I(6.0), I(10.3), I(3.5), color="#3F51B5", width=1.5)
    add_text(slide, I(9.5), I(4.7), I(0.8), I(0.2),
             "SQL", font_size=8, color="#3F51B5", bold=True, align=PP_ALIGN.CENTER)

    # Arrow: MinIO -> Time Travel
    add_arrow(slide, I(4.5), I(6.0), I(10.3), I(4.0), color="#FFA000", width=1.5)
    add_text(slide, I(7.0), I(5.4), I(1.5), I(0.2),
             "USD Download", font_size=7, color="#FFA000", bold=True, align=PP_ALIGN.CENTER)

    # Arrow: Iceberg -> Dashboard
    add_arrow(slide, I(9.6), I(6.5), I(10.3), I(6.5), color="#E91E63", width=1.5)

    # Arrow: Time Travel -> Isaac Sim (restore)
    add_text(slide, I(6.0), I(0.75), I(4.5), I(0.25),
             "Restore: Open root.usda in Isaac Sim",
             font_size=9, color="#00ACC1", bold=True, align=PP_ALIGN.CENTER)
    add_arrow(slide, I(10.3), I(1.5), I(4.8), I(1.2), color="#00ACC1", width=2.5)

    # ══════════════════════════════════════════════════════════
    # Legend
    # ══════════════════════════════════════════════════════════
    legend_items = [
        ("#FFF3E0", "#FF9800", "Isaac Sim"),
        ("#F3E5F5", "#9C27B0", "Nucleus"),
        ("#E8F5E9", "#4CAF50", "Pipeline"),
        ("#FFF8E1", "#FFA000", "MinIO"),
        ("#E8EAF6", "#3F51B5", "Iceberg"),
        ("#E0F7FA", "#00ACC1", "Time Travel"),
        ("#FCE4EC", "#E91E63", "Dashboard"),
    ]

    for i, (fill, border, label) in enumerate(legend_items):
        x = I(0.5 + i * 1.8)
        y = I(7.15)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                      x, y, I(0.25), I(0.2))
        box.fill.solid()
        box.fill.fore_color.rgb = hex_to_rgb(fill)
        box.line.color.rgb = hex_to_rgb(border)
        box.line.width = Pt(1)

        add_text(slide, Inches(0.5 + i * 1.8 + 0.3), y, I(1.4), I(0.2),
                 label, font_size=8, color="#333333")

    # ── Save ──
    out_dir = os.path.dirname(os.path.abspath(__file__))
    out_path = os.path.join(out_dir, "brandt_architecture.pptx")
    prs.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
